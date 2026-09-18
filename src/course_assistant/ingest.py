from __future__ import annotations

import hashlib
import re
from pathlib import Path

from .models import DocumentChunk, SourceEvidence


def _chunk_id(document_name: str, location: str | None, text: str) -> str:
    raw = f"{document_name}|{location or ''}|{text}".encode("utf-8")
    return hashlib.sha1(raw).hexdigest()[:16]


def ingest_text(
    text: str,
    document_name: str,
    section: str | None = None,
    page_or_slide: str | None = None,
    image_path: str | None = None,
) -> list[DocumentChunk]:
    cleaned = re.sub(r"\s+", " ", text).strip()
    if not cleaned:
        return []
    source = SourceEvidence(
        document=document_name,
        page_or_slide=page_or_slide,
        section=section,
        excerpt=cleaned,
        image_path=image_path,
    )
    return [
        DocumentChunk(
            chunk_id=_chunk_id(document_name, page_or_slide, cleaned),
            text=cleaned,
            source=source,
            modality="visual" if image_path else "text",
        )
    ]


def ingest_file(path: str | Path, artifact_dir: str | Path = "artifacts") -> list[DocumentChunk]:
    """Parse a supported course file while retaining page/slide evidence when available."""
    file_path = Path(path)
    suffix = file_path.suffix.casefold()
    if suffix in {".txt", ".md", ".markdown"}:
        return ingest_text(file_path.read_text(encoding="utf-8"), file_path.name)
    if suffix == ".pdf":
        return _ingest_pdf(file_path, Path(artifact_dir))
    if suffix == ".pptx":
        return _ingest_pptx(file_path)
    if suffix == ".docx":
        return _ingest_docx(file_path)
    raise ValueError(f"unsupported file format: {suffix or '(none)'}")


def _ingest_pdf(path: Path, artifact_dir: Path) -> list[DocumentChunk]:
    try:
        import fitz  # type: ignore
    except ImportError as exc:
        raise RuntimeError("PDF support requires PyMuPDF; install the optional dependencies") from exc
    document = fitz.open(path)
    chunks: list[DocumentChunk] = []
    output_dir = artifact_dir / path.stem
    output_dir.mkdir(parents=True, exist_ok=True)
    for page_index, page in enumerate(document, start=1):
        image_path = output_dir / f"page-{page_index}.png"
        pixmap = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
        pixmap.save(image_path)
        chunks.extend(
            ingest_text(
                page.get_text("text"),
                path.name,
                page_or_slide=f"page {page_index}",
                image_path=str(image_path),
            )
        )
    return chunks


def _ingest_pptx(path: Path) -> list[DocumentChunk]:
    try:
        from pptx import Presentation  # type: ignore
    except ImportError as exc:
        raise RuntimeError("PPTX support requires python-pptx; install the optional dependencies") from exc
    presentation = Presentation(path)
    chunks: list[DocumentChunk] = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        text = " ".join(shape.text for shape in slide.shapes if hasattr(shape, "text"))
        chunks.extend(ingest_text(text, path.name, page_or_slide=f"slide {slide_index}"))
    return chunks


def _ingest_docx(path: Path) -> list[DocumentChunk]:
    try:
        from docx import Document  # type: ignore
    except ImportError as exc:
        raise RuntimeError("DOCX support requires python-docx; install the optional dependencies") from exc
    document = Document(path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    return ingest_text(text, path.name)
