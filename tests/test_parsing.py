"""Unit tests for parsing (offline; generates fixtures on the fly)."""
import pytest

from course_assistant import parsing


@pytest.fixture
def work(tmp_path):
    return tmp_path / "pages"


def test_txt_parse(work):
    work.mkdir(parents=True, exist_ok=True)
    src = work / "lecture.txt"
    src.write_text("Welcome to week 2.\nWe cover vector embeddings.", encoding="utf-8")
    doc = parsing.parse_document(src, work)
    assert doc.title == "lecture"
    assert len(doc.pages) == 1
    assert "vector embeddings" in doc.pages[0].text


def test_markdown_parse(work):
    work.mkdir(parents=True, exist_ok=True)
    src = work / "notes.md"
    src.write_text("# Heading\n\nSome **bold** content.", encoding="utf-8")
    doc = parsing.parse_document(src, work)
    assert "Heading" in doc.pages[0].text


def test_pdf_parse(work):
    work.mkdir(parents=True, exist_ok=True)
    import pymupdf
    pdf = work / "deck.pdf"
    doc = pymupdf.open()
    page = doc.new_page(width=400, height=300)
    page.insert_text((40, 60), "Slide one content")
    doc.save(pdf)
    doc.close()
    parsed = parsing.parse_document(pdf, work / "out", parser=None)
    assert parsed.title == "deck"
    assert len(parsed.pages) >= 1
    p0 = parsed.pages[0]
    assert "Slide one content" in p0.text
    assert p0.image_path and p0.image_path.endswith(".png")


def test_pptx_text_extraction(work):
    work.mkdir(parents=True, exist_ok=True)
    from pptx import Presentation
    src = work / "pres.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "My Slide Title"
    prs.save(str(src))
    doc = parsing.parse_document(src, work, allow_libreoffice=False)
    assert doc.pages
    assert "My Slide Title" in doc.pages[0].text


def test_docx_text_extraction(work):
    work.mkdir(parents=True, exist_ok=True)
    from docx import Document
    src = work / "report.docx"
    d = Document()
    d.add_paragraph("First paragraph of the report.")
    d.add_paragraph("A second important claim.")
    d.save(str(src))
    doc = parsing.parse_document(src, work, allow_libreoffice=False)
    assert "First paragraph of the report" in doc.pages[0].text
    assert "second important claim" in doc.pages[0].text


def test_unsupported_format_raises(work):
    work.mkdir(parents=True, exist_ok=True)
    src = work / "data.xlsx"
    src.write_bytes(b"not a supported file")
    with pytest.raises(parsing.UnsupportedFormatError):
        parsing.parse_document(src, work)
