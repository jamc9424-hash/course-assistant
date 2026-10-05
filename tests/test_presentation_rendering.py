import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from course_assistant import app
from course_assistant.ingest import ingest_file


def test_legacy_presentation_requires_libreoffice_when_conversion_is_unavailable(tmp_path, monkeypatch):
    source = tmp_path / "lecture.odp"
    source.write_bytes(b"placeholder")
    monkeypatch.setattr("course_assistant.ingest.shutil.which", lambda _: None)

    with pytest.raises(RuntimeError, match="LibreOffice"):
        ingest_file(source, tmp_path / "artifacts")


def test_converted_legacy_presentation_keeps_original_document_name(tmp_path, monkeypatch):
    source = tmp_path / "lecture.odp"
    source.write_bytes(b"placeholder")
    converted = tmp_path / "lecture.pdf"
    converted.write_bytes(b"converted")
    seen = {}

    monkeypatch.setattr("course_assistant.ingest._convert_presentation", lambda *_: converted)

    def fake_ingest_pdf(path, artifact_dir, location_prefix="page", document_name=None):
        seen.update(path=path, artifact_dir=artifact_dir, prefix=location_prefix, name=document_name)
        return []

    monkeypatch.setattr("course_assistant.ingest._ingest_pdf", fake_ingest_pdf)
    ingest_file(source, tmp_path / "artifacts")

    assert seen["name"] == "lecture.odp"
    assert seen["prefix"] == "slide"


def test_pptx_extracts_embedded_visual_when_libreoffice_is_unavailable(tmp_path, monkeypatch):
    source = tmp_path / "visual-lecture.pptx"
    picture = tmp_path / "meme.png"
    Image.new("RGB", (640, 360), "orange").save(picture)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1)).text = "Vibe Coding on Prod"
    slide.shapes.add_picture(str(picture), Inches(1), Inches(1.5), width=Inches(8))
    presentation.save(source)
    monkeypatch.setattr("course_assistant.ingest._convert_presentation", lambda *_: None)

    chunks = ingest_file(source, tmp_path / "artifacts")

    assert chunks
    assert chunks[0].source.page_or_slide == "slide 1"
    assert chunks[0].source.image_path
    extracted = tmp_path / "artifacts" / "slides" / "slide-1-embedded.png"
    assert extracted.is_file()
    assert chunks[0].source.image_path == str(extracted)


def test_invalid_deployment_port_is_rejected_before_launch(monkeypatch):
    monkeypatch.setenv("PORT", "70000")
    monkeypatch.setattr(app, "build_app", lambda: pytest.fail("app should not launch"))

    with pytest.raises(RuntimeError, match="between 1 and 65535"):
        app.main()
